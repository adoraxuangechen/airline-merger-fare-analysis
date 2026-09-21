#!/usr/bin/env python3
"""Build the editable paper, technical supplement and data note from outputs.

Usage: python documents/build_documents.py [--root /path/to/replication]
Requires python-docx. PDF export is a separate rendering step (see README).
"""
import argparse, csv, json, re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

def load_csv(p):
    with p.open() as f: return list(csv.DictReader(f))

def add_link(p, label, url):
    link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'24556C');rp.append(color)
    run.append(rp);t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);p._p.append(link)

def rich(p,text):
    # Long source URLs are accessible links, not lines of unreadable URL text.
    bits=re.split(r'(https?://[^\s。]+)',text)
    for bit in bits:
        if bit.startswith('http'):
            url=bit.rstrip('.,');add_link(p,'Source',url)
        else:p.add_run(bit)

def table(doc,headers,rows,widths,cn=False):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,w in enumerate(widths):t.columns[j].width=Inches(w)
    for j,h in enumerate(headers):t.rows[0].cells[j].text=h
    for vals in rows:
        cells=t.add_row().cells
        for j,v in enumerate(vals):cells[j].text=str(v)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge,sz in [('top',75),('bottom',75),('left',95),('right',95)]:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(sz));e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'37464F' if i==0 else ('F1F4F5' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.04
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or cn else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size=Pt(9.3 if cn else 10);r.font.bold=(i==0)
                    if i==0:r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def equation(doc):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    om=OxmlElement('m:oMath')
    def run(text):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
    def sub(base,idx):
        s=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(run(base));u=OxmlElement('m:sub');u.append(run(idx));s.append(e);s.append(u);return s
    for item in [sub('y','rt'),run(' = '),sub('α','r'),run(' + '),sub('λ','t'),run(' + β('),sub('T','r'),run(' × '),sub('Post','t'),run(') + '),sub('ε','rt')]:om.append(item)
    p._p.append(om)

def setup(cn):
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.76);s.bottom_margin=Inches(.74);s.left_margin=s.right_margin=Inches(.85)
    s.footer_distance=Inches(.32)
    # Clear decorative borders and theme font overrides in the bundled template.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts
        for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            rf.attrib.pop(qn('w:'+k),None)
        rf.set(qn('w:eastAsia'),'Hiragino Sans GB')
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5 if cn else 11.5)
    normal.paragraph_format.line_spacing=1.13 if cn else 1.08;normal.paragraph_format.space_after=Pt(8 if cn else 8)
    normal.paragraph_format.widow_control=True
    doc.styles['Title'].font.size=Pt(19 if cn else 21);doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(6)
    doc.styles['Subtitle'].font.size=Pt(14);doc.styles['Subtitle'].paragraph_format.space_after=Pt(7)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].font.bold=True
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(5);doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
    doc.styles['Heading 1'].paragraph_format.keep_with_next=True
    doc.styles['Caption'].font.size=Pt(9 if cn else 9.1);doc.styles['Caption'].font.italic=False;doc.styles['Caption'].font.bold=False
    doc.styles['Caption'].paragraph_format.space_after=Pt(10);doc.styles['Caption'].paragraph_format.line_spacing=1.04
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run();r.font.size=Pt(9);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    doc.core_properties.author='Xuange (Adora) Chen';doc.core_properties.subject='Delta–Northwest merger fare analysis'
    return doc

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1]);ap.add_argument('--only',choices=['writing_sample','data_methods','technical_supplement']);args=ap.parse_args();root=args.root
    r=json.loads((root/'results/results.json').read_text());a=json.loads((root/'results/data_audit.json').read_text());o=json.loads((root/'results/original_reproduction.json').read_text())
    base=r['baseline'];co=base['coef']['did'];ev=r['event']['all_pre'];early=r['early_pre_event']['all_pre'];flow=r['sample_flow']
    assert base['N']==6259 and base['treated_routes']==156 and base['control_routes']==105
    assert abs(co['b']-(-.0588984578886))<1e-9
    val={'baseline_pct':f"{co['percent']:.2f}",'baseline_ci':f"[{co['percent_lo']:.2f}%, {co['percent_hi']:.2f}%]",'baseline_b':f"{co['b']:.4f}",'baseline_se':f"{co['se']:.4f}",'treated':base['treated_routes'],'controls':base['control_routes'],'fixed_pct':f"{r['fixed_pre_pax']['coef']['did']['percent']:.2f}",'crisis_pct':f"{r['exclude_crisis_window']['coef']['did']['percent']:.2f}",'any_pct':f"{r['anycode_2007']['coef']['did']['percent']:.2f}",'any_treated':f"{r['anycode_2007']['treated_routes']:,}",'any_controls':f"{r['anycode_2007']['control_routes']:,}",'broad_legacy_pct':f"+{r['broad_reconstruction']['coef']['did']['percent']:.2f}",'broad_all_pct':f"+{r['broad_all_controls']['coef']['did']['percent']:.2f}",'old_event_m8':f"{o['event']['coef']['event_-8']['b']:+.4f}",'old_event_p8':f"{o['event']['coef']['event_8']['b']:+.4f}",'old_event_n':f"{o['event']['N']:,}",'figure_match_note':'在逐点走势和误差棒上相符','distance_note':'主样本的比较路线 SJU-STT 在 24 季的距离字段均为 0，应视为不能作物理距离解释的记录。删除这条路线后票价对比为 '+f"{r['exclude_zero_distance_route']['coef']['did']['percent']:.2f}%"+'，不改变正文判断'}
    for name,test in [('event',ev),('early',early)]:
        for key in ['df_num','df_den']:val[name+('_df1' if key=='df_num' else '_df2')]=test[key]
        val[name+'_f']=f"{test['F']:.3f}";val[name+'_p']=f"{test['p']:.3g}"
    sample_rows=[['Input aggregate records',f"{flow['raw_rows']:,}"],['Records after cleaning',f"{flow['clean_rows']:,}"],['All airport-pair–quarter cells',f"{flow['clean_route_quarters']:,}"],['Routes eligible under 2007 traffic rule',f"{flow['eligible_2007_routes']:,}"],['Main overlap / comparison routes',f"{base['treated_routes']} / {base['control_routes']}"],['Main route-quarter observations',f"{base['N']:,}"]]
    d={z['group']:z for z in r['descriptive']};g1=d['overlap'];g0=d['unexposed_legacy']
    balance_rows=[['Routes',156,105],['Mean fare ($)',f"{g1['mean_2007_fare']:.2f}",f"{g0['mean_2007_fare']:.2f}"],['Quarterly sampled passengers',f"{g1['mean_2007_pax']:,.0f}",f"{g0['mean_2007_pax']:,.0f}"],['Distance field (miles)',f"{g1['mean_distance']:,.0f}",f"{g0['mean_distance']:,.0f}"],['One-coupon passenger share',f"{100*g1['mean_direct_share']:.2f}%",f"{100*g0['mean_direct_share']:.2f}%"]]
    reg_rows=[]
    for key,label in [('baseline','Equal route weights'),('fixed_pre_pax','Fixed 2007 passenger weights'),('balanced','Complete 24-quarter routes'),('omit_announcement_completion','Omit 2008Q2–2008Q4'),('exclude_crisis_window','Omit 2008Q3–2009Q4'),('exclude_zero_distance_route','Omit zero-distance route')]:
        z=r[key];c=z['coef']['did'];reg_rows.append([label,f"{c['b']:.4f}",f"{c['se']:.4f}",f"{c['percent']:.2f}",f"{z['N']:,}"])
    bridge=load_csv(root/'results/specification_bridge.csv');labels=['A 原代码','B 仅修正固定效应','C 改为旅客加权票价','D 加入 2007 样本门槛','E 改用可观测重叠处理组','F 收窄比较组成为主规格']
    bridge_rows=[[labels[i],f"{float(z['percent']):+.2f}%",f"{int(z['treated_routes']):,}",f"{int(z['control_routes']):,}",f"{int(z['N']):,}"] for i,z in enumerate(bridge)]
    fields=[['cr1、cr2','运营承运人代码集合；不是票务承运人或有顺序的两个航段。'],['yr、qtr','记录的年份和季度。'],['cop','0 为一 coupon，1 为两 coupon 类别。'],['ap1、ap2','按字母排序的两个机场，方向合并。'],['pax','该行代表的样本旅客行程计数，用于恢复均价。'],['nsdst','机场对距离字段；0 值不能按真实零距离解释。'],['avprc','该行单程等价平均票价，名义美元。'],['avdst','该行平均行程距离；本次主估计不使用。']]
    files=[['data_audit.json','输入指纹、字段取值、缺失、重复、版本及距离核验。'],['row_cleaning_flow.csv','逐步删去与保留的行数。'],['route_classification.csv','所有 2007 年路线的资格、分类和份额判据。'],['analysis_panel.csv','6,259 个主回归市场季度和票价构造量。'],['regression_results.csv','主结果及所有已运行的 pooled 敏感性估计。'],['event_coefficients.csv','新事件研究系数、标准误、区间和百分比变换。'],['specification_bridge.csv','原始 +5% 到主规格 −5.72% 的逐项比较。'],['original_reproduction.json','忠实重跑原始回归，包括不正确的旧事件规格。'],['results.json','样本量、全部回归、联合检验及数值验证。']]
    tables={'sample':(['Stage','Count'],sample_rows,[5.0,1.8]),'balance':(['Characteristic','Overlap','Comparison'],balance_rows,[3.8,1.5,1.5]),'regression':(['Specification','β','SE','Change %','N'],reg_rows,[3.0,.85,.85,1.05,1.05]),'bridge':(['步骤','相对变化','处理路线','比较路线','市场季度'],bridge_rows,[2.65,.9,1.0,1.05,1.2]),'fields':(['字段','含义与用途'],fields,[1.45,5.35]),'files':(['文件','用途'],files,[2.65,4.15])}
    definition_rows=[]
    for key,label in [('baseline','Material overlap in 2007'),('anycode_2007','Any code in 2007'),('broad_reconstruction','Before completion / legacy controls'),('broad_all_controls','Before completion / all controls')]:
        z=r[key];c=z['coef']['did'];definition_rows.append([label,f"{z['treated_routes']:,}",f"{z['control_routes']:,}",f"{c['percent']:+.2f}"])
    tables['definitions']=(['Route definition','Overlap routes','Comparison routes','Change %'],definition_rows,[3.15,1.25,1.35,1.05])
    labels_en=['A Original implementation','B Correct fixed effects only','C Passenger-weighted fare','D Add 2007 eligibility','E Change overlap classification','F Restrict comparison routes']
    bridge_en=[[labels_en[i],f"{float(z['percent']):+.2f}",f"{int(z['treated_routes']):,}",f"{int(z['control_routes']):,}",f"{int(z['N']):,}"] for i,z in enumerate(bridge)]
    tables['bridge_en']=(['Specification','Change %','Overlap','Comparison','N'],bridge_en,[2.55,.9,.95,1.15,1.25])
    fields_en=[['cr1, cr2','Observed operating-carrier set; segment order is not retained.'],['yr, qtr','Calendar year and quarter.'],['cop','0 identifies one coupon; 1 identifies two coupons.'],['ap1, ap2','Alphabetically ordered airport endpoints; directions combined.'],['pax','Sampled passenger journeys represented by the aggregate record.'],['avprc','Mean one-way-equivalent fare in nominal dollars.'],['nsdst','Airport-pair distance field; zero is not a physical distance estimate.'],['avdst','Mean itinerary distance; unused in the main estimation.']]
    tables['fields_en']=(['Field','Definition'],fields_en,[1.25,5.55])
    event_rows=[]
    ecoefs={int(z['t']):z for z in load_csv(root/'results/event_coefficients.csv')}
    for t in range(1,25):
        quarter=f'{2005+(t-1)//4}Q{(t-1)%4+1}'
        if t==15:event_rows.append([quarter,'0 (reference)','—','—'])
        else:
            z=ecoefs[t];event_rows.append([quarter,f"{float(z['b']):+.4f}",f"{float(z['se']):.4f}",f"[{float(z['lo']):+.4f}, {float(z['hi']):+.4f}]"])
    tables['events']=(['Quarter','Log-fare contrast','SE','95% interval'],event_rows,[1.35,1.8,1.1,2.55])
    val['matching_pct']=f"{r['matched']['coef']['did']['percent']:.2f}"
    val['matching_f']=f"{r['matched_event']['all_pre']['F']:.3f}"
    delivery=root/'delivery';delivery.mkdir(exist_ok=True)
    for name,cn in [('writing_sample',False),('data_methods',True),('technical_supplement',False)]:
        if args.only and name!=args.only:continue
        source_path=root/'documents'/f'{name}.md'
        # The public repository contains the two English templates only.
        if not source_path.exists() and not args.only:continue
        src=source_path.read_text()
        for k,v in val.items():src=src.replace('{{'+k+'}}',str(v))
        assert '{{' not in src, f'Unresolved placeholders: {name}'
        src=src.replace('@page\n','@page\n\n')
        doc=setup(cn)
        pending_break=False
        for block in src.split('\n\n'):
            block=block.strip()
            if not block:continue
            for directive in block.splitlines() if block.startswith('@mono') or block.startswith('# ') else [block]:
                if directive.startswith('# '):doc.add_paragraph(directive[2:],style='Title')
                elif directive.startswith('## '):doc.add_paragraph(directive[3:],style='Subtitle')
                elif directive.startswith('### '):
                    p=doc.add_paragraph(directive[4:],style='Heading 1')
                    if pending_break:p.paragraph_format.page_break_before=True;pending_break=False
                elif directive=='@page':pending_break=True
                elif directive.startswith('@byline '):
                    p=doc.add_paragraph(directive[8:]);p.paragraph_format.space_after=Pt(17)
                    for rr in p.runs:rr.font.size=Pt(10)
                elif directive.startswith('@table '):
                    # Captions may follow directly without a blank line.
                    lines=directive.splitlines();table(doc,*tables[lines[0][7:]],cn=cn)
                    for line in lines[1:]:rich(doc.add_paragraph(style='Caption'),line.removeprefix('@caption '))
                elif directive.startswith('@figure '):
                    lines=directive.splitlines();p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
                    rr=p.add_run();rr.add_picture(str(root/'results'/lines[0][8:]),width=Inches(6.8))
                    for line in lines[1:]:rich(doc.add_paragraph(style='Caption'),line.removeprefix('@caption '))
                elif directive=='@equation':equation(doc)
                elif directive.startswith('@caption '):rich(doc.add_paragraph(style='Caption'),directive[9:])
                elif directive.startswith('@mono '):
                    p=doc.add_paragraph(directive[6:]);p.paragraph_format.space_after=Pt(4)
                    for rr in p.runs:rr.font.name='Courier New';rr.font.size=Pt(9)
                elif directive.startswith('@english '):
                    p=doc.add_paragraph();rich(p,directive[9:])
                    for rr in p.runs:rr.font.size=Pt(10.6)
                else:rich(doc.add_paragraph(),directive)
        doc.core_properties.title={'writing_sample':'Assessing Fare Changes After the Delta–Northwest Merger','technical_supplement':'Delta Northwest Fare Analysis Technical Supplement','data_methods':'Delta Northwest 票价分析的数据处理与核验说明'}[name]
        doc.save(delivery/(name+'.docx'))
        (delivery/(name+'.md')).write_text(src)
        print(f'Built {name}.docx')
    (delivery/'document_values.json').write_text(json.dumps(val,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
